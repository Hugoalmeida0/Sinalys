"""Leitura de planilhas (.xlsx, .xls, .csv) no formato "uma linha = um dicionário por cabeçalho"."""

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime

EXTENSOES = (".xlsx", ".xls", ".csv")
REGEX_NUMERO = re.compile(r"^-?\d+(\.\d+)?$")


@dataclass
class Aba:
    nome: str
    linhas: list[dict]
    colunas: list[str]


def detectar_tipo_origem(nome_arquivo: str) -> str:
    nome = nome_arquivo.lower()
    if nome.endswith(".csv"):
        return "csv"
    if nome.endswith(".xlsx") or nome.endswith(".xls"):
        return "excel"
    raise ValueError("Extensão de arquivo não suportada. Use .xlsx, .xls ou .csv.")


def _cabecalhos(primeira: list) -> list[str]:
    vistos: dict[str, int] = {}
    cabecalhos = []
    for i, valor in enumerate(primeira):
        nome = str(valor).strip() if valor not in (None, "") else f"__EMPTY_{i}"
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome}_{vistos[nome]}"
        else:
            vistos[nome] = 0
        cabecalhos.append(nome)
    return cabecalhos


def _montar_aba(nome: str, matriz: list[list]) -> Aba:
    matriz = [linha for linha in matriz if any(c not in (None, "") for c in linha)]
    if not matriz:
        return Aba(nome, [], [])
    colunas = _cabecalhos(matriz[0])
    linhas = []
    for bruta in matriz[1:]:
        linhas.append({coluna: (bruta[i] if i < len(bruta) and bruta[i] != "" else None) for i, coluna in enumerate(colunas)})
    return Aba(nome, linhas, colunas)


def _converter_csv(valor: str):
    """Números no CSV viram número, como faz o leitor de planilhas no navegador."""
    texto = valor.strip()
    if REGEX_NUMERO.match(texto) and not (len(texto) > 1 and texto.startswith("0") and "." not in texto):
        return float(texto) if "." in texto else int(texto)
    return valor


def ler_planilha(conteudo: bytes, tipo_origem: str, nome_arquivo: str = "") -> list[Aba]:
    if tipo_origem == "csv":
        texto = conteudo.decode("utf-8-sig", errors="replace")
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t") if texto.strip() else csv.excel
        matriz = [[_converter_csv(c) for c in linha] for linha in csv.reader(io.StringIO(texto), dialeto)]
        return [_montar_aba("", matriz)]

    if nome_arquivo.lower().endswith(".xls"):
        import xlrd

        livro = xlrd.open_workbook(file_contents=conteudo)
        abas = []
        for folha in livro.sheets():
            matriz = []
            for r in range(folha.nrows):
                linha = []
                for c in range(folha.ncols):
                    celula = folha.cell(r, c)
                    if celula.ctype == xlrd.XL_CELL_DATE:
                        linha.append(xlrd.xldate.xldate_as_datetime(celula.value, livro.datemode))
                    else:
                        linha.append(celula.value)
                matriz.append(linha)
            abas.append(_montar_aba(folha.name, matriz))
        return abas

    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    return [_montar_aba(folha.title, [list(linha) for linha in folha.iter_rows(values_only=True)]) for folha in livro.worksheets]


def valor_como_texto(valor) -> str:
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime):
        return valor.isoformat()
    return str(valor)
